"""Excel generation using openpyxl."""
from io import BytesIO
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def _header_fill():
    return PatternFill("solid", fgColor="1F4E79")


def _border():
    s = Side(border_style="thin", color="888888")
    return Border(left=s, right=s, top=s, bottom=s)


def generate_excel_report(title: str, columns: list[str], rows: list[list],
                          sheet_name: str = "Report", summary: dict | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:30]

    row_idx = 1
    ws.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=max(len(columns), 4))
    c = ws.cell(row=row_idx, column=1, value=title)
    c.font = Font(size=14, bold=True, color="1F4E79")
    c.alignment = Alignment(horizontal="center")
    row_idx += 1

    if summary:
        for k, v in summary.items():
            ws.cell(row=row_idx, column=1, value=k).font = Font(bold=True)
            ws.cell(row=row_idx, column=2, value=str(v))
            row_idx += 1
        row_idx += 1

    header_row = row_idx
    for col_i, col_name in enumerate(columns, start=1):
        c = ws.cell(row=header_row, column=col_i, value=col_name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = _header_fill()
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = _border()
    row_idx += 1

    for row in rows:
        for col_i, val in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_i, value=val if not isinstance(val, Decimal) else float(val))
            cell.border = _border()
            if isinstance(val, (int, float, Decimal)):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right")
        row_idx += 1

    # Auto-filter + freeze
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{row_idx - 1}"
    ws.freeze_panes = f"A{header_row + 1}"

    # Column widths
    for col_i, col_name in enumerate(columns, start=1):
        width = max(12, min(40, len(str(col_name)) + 4))
        ws.column_dimensions[get_column_letter(col_i)].width = width

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_payslip_excel(employee, record, payroll_month, org) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Payslip"
    ws.merge_cells("A1:D1")
    ws["A1"] = org.name if org else "Homoeopathic Medical College"
    ws["A1"].font = Font(size=14, bold=True, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="center")

    months = ["", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    ws.merge_cells("A2:D2")
    ws["A2"] = f"Payslip — {months[payroll_month.month]} {payroll_month.year}"
    ws["A2"].font = Font(size=12, bold=True)
    ws["A2"].alignment = Alignment(horizontal="center")

    info = [
        ("Employee Code", employee.emp_code, "DOJ", str(employee.date_of_joining or "")),
        ("Name", f"{employee.first_name} {employee.last_name or ''}".strip(),
         "Department", employee.department.name if employee.department else ""),
        ("Designation", employee.designation.name if employee.designation else "",
         "UAN", employee.uan or ""),
        ("Working Days", record.working_days, "Paid Days", float(record.paid_days)),
    ]
    r = 4
    for row in info:
        for i, val in enumerate(row):
            cell = ws.cell(row=r, column=i + 1, value=val)
            if i % 2 == 0:
                cell.font = Font(bold=True)
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Earnings").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=r, column=1).fill = _header_fill()
    ws.cell(row=r, column=2, value="Amount (₹)").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=r, column=2).fill = _header_fill()
    ws.cell(row=r, column=3, value="Deductions").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=r, column=3).fill = _header_fill()
    ws.cell(row=r, column=4, value="Amount (₹)").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=r, column=4).fill = _header_fill()
    r += 1
    earnings = [
        ("Basic", record.basic), ("DA", record.da), ("HRA", record.hra),
        ("Medical", record.medical_allowance), ("Conveyance", record.conveyance_allowance),
        ("Transport", record.transport_allowance), ("Special", record.special_allowance),
        ("Other", record.other_allowance), ("CEA", record.children_education_allowance),
        ("Uniform", record.uniform_allowance), ("Telephone", record.telephone_allowance),
        ("Internet", record.internet_allowance), ("Research", record.research_allowance),
        ("Overtime", record.overtime_pay), ("Bonus", record.bonus),
        ("Arrears", record.arrears), ("Reimb.", record.reimbursement),
    ]
    deductions = [
        ("EPF", record.epf_employee), ("ESI", record.esi_employee),
        ("TDS", record.tds), ("PT", record.professional_tax),
        ("LWF", record.lwf_employee), ("Loan EMI", record.loan_emi),
        ("Advance", record.advance_recovery), ("Other", record.other_deductions),
    ]
    rows = max(len(earnings), len(deductions))
    for i in range(rows):
        if i < len(earnings):
            ws.cell(row=r + i, column=1, value=earnings[i][0])
            ws.cell(row=r + i, column=2, value=float(earnings[i][1] or 0)).number_format = "#,##0.00"
        if i < len(deductions):
            ws.cell(row=r + i, column=3, value=deductions[i][0])
            ws.cell(row=r + i, column=4, value=float(deductions[i][1] or 0)).number_format = "#,##0.00"
    r += rows
    ws.cell(row=r, column=1, value="GROSS").font = Font(bold=True)
    ws.cell(row=r, column=2, value=float(record.gross_salary)).font = Font(bold=True)
    ws.cell(row=r, column=2).number_format = "#,##0.00"
    ws.cell(row=r, column=3, value="TOTAL DEDUCTIONS").font = Font(bold=True)
    ws.cell(row=r, column=4, value=float(record.total_deductions)).font = Font(bold=True)
    ws.cell(row=r, column=4).number_format = "#,##0.00"
    r += 2
    ws.cell(row=r, column=1, value="NET PAY").font = Font(bold=True, size=14)
    ws.cell(row=r, column=2, value=float(record.net_salary)).font = Font(bold=True, size=14)
    ws.cell(row=r, column=2).number_format = "#,##0.00"
    ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor="C6E0B4")

    for c in "ABCD":
        ws.column_dimensions[c].width = 22

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
