"""Excel generation and bulk-upload parsing service."""
import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, List, Optional, Tuple

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import Session


# ──────────────────────────────────────────
# Style helpers
# ──────────────────────────────────────────

_HEADER_FILL = PatternFill(start_color="1A237E", end_color="1A237E", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_ALT_FILL = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")
_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _apply_header(ws, headers: List[str], row: int = 1):
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _THIN_BORDER
    ws.row_dimensions[row].height = 20


def _auto_width(ws, min_width=10, max_width=50):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                val = str(cell.value) if cell.value is not None else ""
                max_len = max(max_len, len(val))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = max(min_width, min(max_width, max_len + 2))


# ──────────────────────────────────────────
# Generic report
# ──────────────────────────────────────────

def generate_report_excel(
    title: str,
    headers: List[str],
    rows: List[List[Any]],
    filters: str = "",
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    # Title row
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 1))
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    if filters:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(len(headers), 1))
        filters_cell = ws.cell(row=2, column=1, value=filters)
        filters_cell.font = Font(italic=True, size=9, color="666666")
        filters_cell.alignment = Alignment(horizontal="center")

    header_row = 3 if filters else 2
    _apply_header(ws, headers, row=header_row)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{header_row}"

    for r_idx, row in enumerate(rows):
        actual_row = header_row + 1 + r_idx
        fill = _ALT_FILL if r_idx % 2 == 1 else None
        for c_idx, value in enumerate(row):
            cell = ws.cell(row=actual_row, column=c_idx + 1, value=value)
            cell.border = _THIN_BORDER
            if fill:
                cell.fill = fill
            # Right-align numbers
            if isinstance(value, (int, float, Decimal)):
                cell.alignment = Alignment(horizontal="right")

    _auto_width(ws)

    # Footer
    footer_row = header_row + 1 + len(rows) + 1
    ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=max(len(headers), 1))
    ws.cell(row=footer_row, column=1, value=f"Generated: {datetime.now().strftime('%d %b %Y %H:%M')} | Rows: {len(rows)}").font = Font(italic=True, size=8, color="888888")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ──────────────────────────────────────────
# Bulk upload template
# ──────────────────────────────────────────

BULK_HEADERS = [
    "admission_number*",
    "student_name*",
    "father_name",
    "course_name",
    "batch_name",
    "session_name",
    "professional_year",
    "roll_number",
    "receipt_date*(DD/MM/YYYY)",
    "fee_head_1_name*",
    "fee_head_1_amount*",
    "fee_head_2_name",
    "fee_head_2_amount",
    "fee_head_3_name",
    "fee_head_3_amount",
    "payment_mode_1*",
    "payment_amount_1*",
    "reference_1",
    "bank_name_1",
    "payment_mode_2",
    "payment_amount_2",
    "reference_2",
    "remarks",
]

SAMPLE_ROW = [
    "HMC-FEE-001",
    "Student Name",
    "Father Name",
    "BHMS",
    "2024-2029",
    "2025-2026",
    "1",
    "101",
    "01/04/2025",
    "Tuition Fee",
    "50000.00",
    "Development Fee",
    "10000.00",
    "",
    "",
    "CASH",
    "60000.00",
    "",
    "",
    "",
    "",
    "",
    "Sample receipt",
]

VALID_MODES = ["CASH", "BANK", "CHEQUE", "DD", "UPI", "CARD", "OTHER"]


def generate_bulk_upload_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Receipts"

    _apply_header(ws, BULK_HEADERS, row=1)

    # Sample row
    for c_idx, val in enumerate(SAMPLE_ROW, 1):
        cell = ws.cell(row=2, column=c_idx, value=val)
        cell.fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")

    # Data validation for payment modes
    mode_formula = f'"{",".join(VALID_MODES)}"'
    dv1 = DataValidation(type="list", formula1=mode_formula, allow_blank=True, showDropDown=False)
    dv2 = DataValidation(type="list", formula1=mode_formula, allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv1)
    ws.add_data_validation(dv2)
    dv1.add(f"P2:P1000")  # payment_mode_1
    dv2.add(f"T2:T1000")  # payment_mode_2

    # Instructions sheet
    ws2 = wb.create_sheet("Instructions")
    instructions = [
        ("Instructions for Bulk Receipt Upload", True),
        ("", False),
        ("1. Do NOT change column headers.", False),
        ("2. Fields marked with * are mandatory.", False),
        ("3. Date format: DD/MM/YYYY", False),
        ("4. Payment modes: CASH, BANK, CHEQUE, DD, UPI, CARD, OTHER", False),
        ("5. You can add up to 3 fee heads and 2 payment modes per row.", False),
        ("6. Amounts should be numeric (e.g. 50000.00)", False),
        ("7. Duplicate admission numbers in the same upload are flagged.", False),
        ("8. Upload will be validated before committing. Review errors before confirming.", False),
    ]
    for r_idx, (text, bold) in enumerate(instructions, 1):
        cell = ws2.cell(row=r_idx, column=1, value=text)
        if bold:
            cell.font = Font(bold=True, size=13)
    ws2.column_dimensions["A"].width = 80

    _auto_width(ws)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ──────────────────────────────────────────
# Bulk upload parsing
# ──────────────────────────────────────────

def _parse_date(val: Any) -> Optional[date]:
    if val is None:
        return None
    if isinstance(val, (date, datetime)):
        return val.date() if isinstance(val, datetime) else val
    s = str(val).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _parse_decimal(val: Any) -> Optional[Decimal]:
    if val is None or val == "":
        return None
    try:
        return Decimal(str(val).strip().replace(",", ""))
    except InvalidOperation:
        return None


def parse_bulk_upload_excel(
    file_bytes: bytes, db: Session
) -> Tuple[List[dict], List[dict]]:
    """
    Parse Excel bulk upload file.
    Returns (valid_rows, error_rows).
    Each row is a dict with keys: row_number, raw_data, parsed_data, status, error_message.
    """
    from app.models.student import Student
    from app.models.org import FinancialYear
    from sqlalchemy import select

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as e:
        raise ValueError(f"Cannot open Excel file: {e}")

    ws = wb.active
    rows_iter = list(ws.iter_rows(values_only=True))
    if not rows_iter:
        raise ValueError("Empty spreadsheet.")

    # Detect header row
    header_row_idx = 0
    headers_lower = [str(h).lower().strip().replace("*", "").replace(" ", "_").split("(")[0]
                     for h in (rows_iter[0] or [])]

    valid_rows = []
    error_rows = []
    seen_admission = set()

    for row_idx, raw_row in enumerate(rows_iter[1:], start=2):
        raw = dict(zip(headers_lower, raw_row))
        errors = []

        # Required fields
        adm_no = str(raw.get("admission_number", "") or "").strip()
        student_name = str(raw.get("student_name", "") or "").strip()
        receipt_date_raw = raw.get("receipt_date_(dd/mm/yyyy)") or raw.get("receipt_date")
        receipt_date = _parse_date(receipt_date_raw)
        mode1 = str(raw.get("payment_mode_1", "") or "").strip().upper()
        amount1_raw = raw.get("payment_amount_1")
        amount1 = _parse_decimal(amount1_raw)

        fh1_name = str(raw.get("fee_head_1_name", "") or "").strip()
        fh1_amount = _parse_decimal(raw.get("fee_head_1_amount"))

        if not adm_no:
            errors.append("admission_number is required")
        if not student_name:
            errors.append("student_name is required")
        if not receipt_date:
            errors.append("receipt_date is required and must be DD/MM/YYYY")
        if not fh1_name:
            errors.append("fee_head_1_name is required")
        if fh1_amount is None:
            errors.append("fee_head_1_amount is required and must be numeric")
        if not mode1:
            errors.append("payment_mode_1 is required")
        elif mode1 not in VALID_MODES:
            errors.append(f"payment_mode_1 must be one of {VALID_MODES}")
        if amount1 is None:
            errors.append("payment_amount_1 is required and must be numeric")

        # Duplicate check within file
        if adm_no and adm_no in seen_admission:
            errors.append(f"Duplicate admission_number '{adm_no}' in this upload")
        elif adm_no:
            seen_admission.add(adm_no)

        # Build items
        items = []
        if fh1_name and fh1_amount is not None:
            items.append({"fee_head_name": fh1_name, "amount": str(fh1_amount)})
        fh2_name = str(raw.get("fee_head_2_name", "") or "").strip()
        fh2_amount = _parse_decimal(raw.get("fee_head_2_amount"))
        if fh2_name and fh2_amount is not None:
            items.append({"fee_head_name": fh2_name, "amount": str(fh2_amount)})
        fh3_name = str(raw.get("fee_head_3_name", "") or "").strip()
        fh3_amount = _parse_decimal(raw.get("fee_head_3_amount"))
        if fh3_name and fh3_amount is not None:
            items.append({"fee_head_name": fh3_name, "amount": str(fh3_amount)})

        # Build payments
        payments = []
        if mode1 and amount1 is not None:
            payments.append({
                "payment_mode_name": mode1,
                "amount": str(amount1),
                "reference_number": str(raw.get("reference_1", "") or "").strip() or None,
                "bank_name": str(raw.get("bank_name_1", "") or "").strip() or None,
                "payment_status": "CLEARED",
            })
        mode2 = str(raw.get("payment_mode_2", "") or "").strip().upper()
        amount2 = _parse_decimal(raw.get("payment_amount_2"))
        if mode2 and amount2 is not None:
            payments.append({
                "payment_mode_name": mode2,
                "amount": str(amount2),
                "reference_number": str(raw.get("reference_2", "") or "").strip() or None,
                "payment_status": "CLEARED",
            })

        raw_data = {k: str(v) if v is not None else None for k, v in raw.items()}
        parsed = {
            "admission_number": adm_no,
            "student_name": student_name,
            "father_name": str(raw.get("father_name", "") or "").strip() or None,
            "course_name": str(raw.get("course_name", "") or "").strip() or None,
            "batch_name": str(raw.get("batch_name", "") or "").strip() or None,
            "session_name": str(raw.get("session_name", "") or "").strip() or None,
            "professional_year": int(raw.get("professional_year") or 0) or None,
            "roll_number": str(raw.get("roll_number", "") or "").strip() or None,
            "receipt_date": receipt_date.isoformat() if receipt_date else None,
            "items": items,
            "payments": payments,
            "remarks": str(raw.get("remarks", "") or "").strip() or None,
        }

        row_info = {
            "row_number": row_idx,
            "raw_data": raw_data,
            "parsed_data": parsed,
        }

        if errors:
            row_info["status"] = "FAILED"
            row_info["error_message"] = "; ".join(errors)
            error_rows.append(row_info)
        else:
            row_info["status"] = "SUCCESS"
            row_info["error_message"] = None
            valid_rows.append(row_info)

    return valid_rows, error_rows
